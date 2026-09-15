"""Generate the thesis DOCX and technical figures from the accompanying Markdown."""
from pathlib import Path
import math
import re
import sys
from copy import deepcopy
from tipografija import segments

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / '.tmp_doc_tools'))
import pymupdf
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.text.run import Run

ASSETS = OUT / 'assets'
ASSETS.mkdir(exist_ok=True)
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'
INK = '#24374b'
BLUE = '#eaf0f7'

def figure(name, height, painter):
    im = Image.new('RGB', (1600, height), 'white')
    draw = ImageDraw.Draw(im)
    painter(draw)
    im.save(ASSETS / name, dpi=(240, 240))

def label(d, xy, text, size=29, bold=False, anchor='mm', fill=INK):
    regular=ImageFont.truetype(BOLD if bold else FONT,size)
    italic=ImageFont.truetype('C:/Windows/Fonts/arialbi.ttf' if bold else 'C:/Windows/Fonts/ariali.ttf',size)
    lines=text.split('\n');top=xy[1]-(len(lines)*(size+9)-9)/2
    for idx,line in enumerate(lines):
        runs=[(part,italic if english else regular) for part,english in segments(line)]
        x=xy[0]-sum(font.getlength(part) for part,font in runs)/2
        for part,font in runs:
            d.text((x,top+idx*(size+9)+size),part,font=font,fill=fill,anchor='ls')
            x+=font.getlength(part)

def box(d, rect, title, detail='', dashed=False):
    x0,y0,x1,y1=rect
    d.rounded_rectangle(rect, radius=14, fill=BLUE if not dashed else '#f6f6f6', outline=INK, width=3)
    label(d, ((x0+x1)/2,(y0+y1)/2-(23 if detail else 0)), title, 30, True)
    if detail: label(d, ((x0+x1)/2,(y0+y1)/2+28), detail, 23)

def arrow(d, start, end, text=None, offset=(0,-24)):
    d.line([start,end],fill=INK,width=4)
    a=math.atan2(end[1]-start[1],end[0]-start[0])
    pts=[end]+[(end[0]-18*math.cos(a+s),end[1]-18*math.sin(a+s)) for s in [-.45,.45]]
    d.polygon(pts,fill=INK)
    if text: label(d,((start[0]+end[0])/2+offset[0],(start[1]+end[1])/2+offset[1]),text,22)

def architecture(d):
    box(d,(35,40,455,165),'ESP32 сензори','предвиђени физички слој',True)
    box(d,(35,210,455,335),'Симулатор','реализован извор порука')
    box(d,(620,120,1010,250),'Mosquitto','MQTT посредник')
    arrow(d,(455,105),(620,155),'MQTT')
    arrow(d,(455,270),(620,215),'MQTT',(0,30))
    box(d,(1160,120,1570,250),'Worker','пријем и обрада')
    arrow(d,(1010,185),(1160,185),'JSON')
    box(d,(1160,400,1570,535),'PostgreSQL','опажања, сесије, аларми')
    arrow(d,(1365,250),(1365,400),'упис',(75,0))
    box(d,(620,400,1010,535),'ASP.NET Core API','REST и SignalR')
    arrow(d,(1160,468),(1010,468),'читање',(0,-25))
    box(d,(35,400,455,535),'Angular','наставнички интерфејс')
    arrow(d,(620,465),(455,465),'HTTP / догађаји',(0,-27))
    d.line([(1220,250),(1220,335),(815,335),(815,400)],fill=INK,width=4)
    arrow(d,(815,365),(815,400))
    label(d,(1015,310),'објављивање догађаја',23)
    label(d,(800,600),'API и Worker користе заједничке слојеве Application, Domain и Infrastructure.',24)

def processing(d):
    steps=[('JSON порука','MQTT или HTTP'),('Провера дупликата','eventId / external_id'),('Хеш и опажање','проналажење уређаја'),('Процена ризика','дозвола + правила'),('Упис и обавештење','deviceUpdated; условни аларм')]
    for i,(title,detail) in enumerate(steps):
        y=30+i*143
        box(d,(70,y,770,y+106),title,detail)
        if i<4: arrow(d,(420,y+106),(420,y+143))
    box(d,(920,135,1530,260),'Захтев за положаје','GET /positions')
    box(d,(920,335,1530,460),'Избор очитавања','временски прозор по сесији')
    box(d,(920,535,1530,660),'Процена и изглађивање','WLS / тежиште → Калман')
    arrow(d,(1225,260),(1225,335))
    arrow(d,(1225,460),(1225,535))
    label(d,(1225,740),'Положаји се рачунају при читању.',25)

def sensors(d):
    left,top,right,bottom=250,120,1350,780
    d.rectangle((left,top,right,bottom),outline=INK,width=4)
    for i in range(1,8):
        x=left+i*(right-left)/8
        d.line((x,top,x,bottom),fill='#e4e8ed',width=2)
    for i in range(1,6):
        y=top+i*(bottom-top)/6
        d.line((left,y,right,y),fill='#e4e8ed',width=2)
    pts=[(left,top),(right,top),((left+right)/2,bottom)]
    d.line([pts[0],pts[2],pts[1],pts[0]],fill='#9badbf',width=4)
    for p in pts: d.ellipse((p[0]-14,p[1]-14,p[0]+14,p[1]+14),fill=INK)
    label(d,(left,70),'S1 (0, 0)',28,True)
    label(d,(right,70),'S2 (8, 0)',28,True)
    label(d,(800,835),'S3 (4, 6)',28,True)
    label(d,(800,900),'8 m',27)
    label(d,(1450,450),'6 m',27)
    label(d,(800,970),'Шематски приказ конфигурације; без експерименталних мерења.',24)

def session(d):
    for r,t,s in [((50,70,470,225),'planned','планирана'),((590,70,1010,225),'active','активна'),((1130,70,1550,225),'completed','завршена')]: box(d,r,t,s)
    arrow(d,(470,148),(590,148),'start')
    arrow(d,(1010,148),(1130,148),'stop')
    label(d,(800,295),'Завршена сесија се не покреће поново.',26)

figure('arhitektura.png',650,architecture)
figure('tok_obrade.png',800,processing)
figure('senzori.png',1010,sensors)
figure('sesija.png',350,session)

# Render the faculty emblem from the user-provided template PDF.
ref = pymupdf.open(ROOT / 'Zavrsni_rad_Mirko_Mihajlovic_62_2021_Final.pdf')
ref[0].get_pixmap(matrix=pymupdf.Matrix(4,4), clip=pymupdf.Rect(252,75,363,186)).save(ASSETS / 'grb_pmf.png')

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.3)
sec.left_margin=Cm(3); sec.right_margin=Cm(2.5)
sec.header_distance=Cm(1.2);sec.footer_distance=Cm(1.2)
sec.different_first_page_header_footer=True
normal=doc.styles['Normal']
normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.widow_control=True
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'sr-Cyrl-RS');normal.element.get_or_add_rPr().append(lang)
for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(14);s.paragraph_format.space_after=Pt(9)
    s.paragraph_format.keep_with_next=True
    if name=='Heading 1':s.paragraph_format.page_break_before=True
for name in ['Caption','TOC 1','TOC 2']:
    if name not in doc.styles: doc.styles.add_style(name,1)
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(10 if name=='Caption' else 11)
    s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.line_spacing=1
    s.paragraph_format.space_after=Pt(5)
doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER

def field(p,code):
    run=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=code
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    run._r.extend([begin,ins,end])

foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
field(foot,' PAGE ')

def center(text,size=12,bold=False,before=0,after=6):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
    r=p.add_run(text);r.bold=bold;r.font.size=Pt(size)
    return p

p=center('',after=8);p.add_run().add_picture(str(ASSETS/'grb_pmf.png'),width=Cm(3.4))
center('ИНСТИТУТ ЗА МАТЕМАТИКУ И ИНФОРМАТИКУ\nПРИРОДНО-МАТЕМАТИЧКИ ФАКУЛТЕТ\nУНИВЕРЗИТЕТ У КРАГУЈЕВЦУ',12,after=25)
center('ЗАВРШНИ РАД',14,before=12,after=16)
p=center('СИСТЕМ ЗА ДЕТЕКЦИЈУ БЕЖИЧНИХ\nУРЕЂАЈА ТОКОМ ИСПИТНИХ СЕСИЈА\nПРИМЕНОМ ESP32 ПЛАТФОРМЕ',18,True,after=0)
p.paragraph_format.line_spacing=1.1
borders=OxmlElement('w:pBdr')
for side in ['top','bottom']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'10');el.set(qn('w:space'),'10');borders.append(el)
p._p.get_or_add_pPr().append(borders)
center('',before=105,after=0)
for t in ['Ментор\tСтудент','др Ана Капларевић-Малишић\tАндрија Костовић, 51/2021']:
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(9))
    p.paragraph_format.space_after=Pt(0)
    r=p.add_run(t);r.font.size=Pt(11);r.bold=True
center('Крагујевац, септембар 2026.',12,before=24)

doc.add_page_break()
p=doc.add_paragraph('Садржај');p.runs[0].bold=True;p.runs[0].font.size=Pt(16)
p=doc.add_paragraph();r=p.add_run()
begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' TOC \\o "1-2" \\h \\z \\u '
sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
r._r.extend([begin,ins,sep])
heading_lines=[s for s in (OUT/'Zavrsni_rad_ESP32.md').read_text(encoding='utf-8').splitlines() if s.startswith('#')]
for idx,s in enumerate(heading_lines):
    level=len(s)-len(s.lstrip('#'))
    if level>2:continue
    p=doc.add_paragraph(style='TOC '+str(level));p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if level==2:p.paragraph_format.left_indent=Cm(.5)
    h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),'section_'+str(idx))
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=s[level:].strip();r.append(t);h.append(r);p._p.append(h)
end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');doc.add_paragraph().add_run()._r.append(end)

def add_hyperlink(p,text,url):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),'244C77');pr.append(c)
    r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def add_text(p,text):
    pos=0
    for m in re.finditer(r'https?://[^\s]+',text):
        p.add_run(text[pos:m.start()]);add_hyperlink(p,m.group(),m.group());pos=m.end()
    p.add_run(text[pos:])

lines=(OUT/'Zavrsni_rad_ESP32.md').read_text(encoding='utf-8').replace('\u00ad','').splitlines()
i=0
heading_index=0
picture_number=4
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('{{SLIKA|'):
        _,key,caption,instruction=line[2:-2].split('|',3);picture_number+=1
        p=doc.add_paragraph('МЕСТО ЗА СЛИКУ. '+instruction)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(95);p.paragraph_format.space_after=Pt(95)
        p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
        borders=OxmlElement('w:pBdr')
        for side in ('top','left','bottom','right'):
            edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'9CA6B0');borders.append(edge)
        p._p.get_or_add_pPr().append(borders)
        doc.add_paragraph(f'Слика {picture_number}. {caption}',style='Caption');i+=1;continue
    if line.startswith('#'):
        n=len(line)-len(line.lstrip('#'));p=doc.add_heading(line[n:].strip(),level=n)
        if line.startswith('## 10.3.'):p.paragraph_format.page_break_before=True
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(heading_index));start.set(qn('w:name'),'section_'+str(heading_index))
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(heading_index))
        p._p.insert(1 if p._p.pPr is not None else 0,start);p._p.append(end);heading_index+=1;i+=1;continue
    if line.startswith('!['):
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(OUT/m[2]),width=Cm(15.3))
        doc.add_paragraph(m[1],style='Caption');i+=1;continue
    if line.startswith('```'):
        block=[];i+=1
        while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
        p=doc.add_paragraph();p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(10)
        p.paragraph_format.keep_together=True
        p.paragraph_format.left_indent=Cm(.2);p.paragraph_format.right_indent=Cm(.1)
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F3F5');p._p.get_or_add_pPr().append(shade)
        r=p.add_run('\n'.join(block));r.font.name='Consolas';r.font.size=Pt(9)
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[s.strip() for s in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',s) for s in row):rows.append(row)
            i+=1
        t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid';t.autofit=False
        widths=([7.8,7.7] if len(rows[0])==2 else [4,5.7,5.8] if len(rows[0])==3 else [2,4.2,4,2.5,2.8])
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for ri,row in enumerate(rows):
            cells=t.add_row().cells
            pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if ri==0:
                el=OxmlElement('w:tblHeader');pr.append(el)
            for ci,(c,txt) in enumerate(zip(cells,row)):
                c.width=Cm(widths[ci]);p=c.paragraphs[0]
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                r=p.add_run(txt);r.font.size=Pt(10);r.bold=ri==0
                if ri==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EDF2');c._tc.get_or_add_tcPr().append(sh)
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
        continue
    if re.match(r'Табела \d+:',line):
        p=doc.add_paragraph(line,style='Caption');p.paragraph_format.keep_with_next=True;i+=1;continue
    if line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet');add_text(p,line[2:])
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.keep_together=True
        i+=1;continue
    p=doc.add_paragraph();add_text(p,line)
    if line.startswith('[') and re.match(r'\[\d+\]',line):
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:r.font.size=Pt(11)
    i+=1

settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.core_properties.title='Систем за детекцију бежичних уређаја током испитних сесија применом ESP32 платформе'
doc.core_properties.author='Андрија Костовић'
doc.core_properties.subject='Завршни рад. ESP32, MQTT, .NET и Angular'
doc.core_properties.keywords='ESP32, RSSI, MQTT, локализација, испитне сесије'
doc.core_properties.comments='Технички опис према доступном коду. Биографски подаци за допуну су означени.'
target=OUT/'Zavrsni_rad_Andrija_Kostovic_51_2021_ESP32.docx'
# Apply italic character formatting consistently, including linked contents,
# hyperlinks, headings, code examples and cover text. Preserve all other styles.
for original in list(doc.element.iter(qn('w:r'))):
    if original.find(qn('w:t')) is None or original.find(qn('w:fldChar')) is not None:
        continue
    value=Run(original,None).text
    if not any(english for _,english in segments(value)):
        continue
    parent=original.getparent();index=parent.index(original)
    for offset,(part,english) in enumerate(segments(value)):
        element=OxmlElement('w:r')
        if original.rPr is not None:element.append(deepcopy(original.rPr))
        run=Run(element,None);run.text=part
        if english:run.italic=True
        parent.insert(index+offset,element)
    parent.remove(original)
doc.save(target)
print(target)
