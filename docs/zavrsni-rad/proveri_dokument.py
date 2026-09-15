from pathlib import Path
import json,re,sys
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(OUT.parents[1]/'.tmp_doc_tools'))
import pymupdf
from docx import Document
from docx.oxml.ns import qn

doc=Document(OUT/'Zavrsni_rad_Andrija_Kostovic_51_2021_ESP32.docx')
pdf=pymupdf.open(OUT/'Zavrsni_rad_Andrija_Kostovic_51_2021_ESP32_Prosireni.pdf')
bad_docx=[];count=0
for r in doc.element.iter(qn('w:r')):
    text=''.join(t.text or '' for t in r.iter(qn('w:t')))
    if re.search('[A-Za-z]',text):
        count+=1
        italic=r.find('w:rPr/w:i',namespaces=r.nsmap)
        if italic is None or italic.get(qn('w:val')) in ('0','false'):bad_docx.append(text)
bad_pdf=[]
for idx,page in enumerate(pdf):
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines',[]):
            for span in line['spans']:
                if re.search('[A-Za-z]',span['text']) and not(span['flags']&2):bad_pdf.append([idx+1,span['text']])
source=(OUT/'Zavrsni_rad_ESP32.md').read_text(encoding='utf-8')
normalize=lambda s:re.sub(r'\s+','',s)
corpus=normalize(' '.join(page.get_text(clip=pymupdf.Rect(0,0,page.rect.width,790)) for page in pdf))
prose=[line.removeprefix('- ') for line in source.splitlines() if (len(line)>140 or line.startswith('- ')) and not line.startswith(('#','!','[','{{SLIKA|'))]
missing=[p[:120] for p in prose if normalize(p) not in corpus]
report={
    'pdf_pages':len(pdf),'pdf_bookmarks':len(pdf.get_toc()),
    'docx_tables':len(doc.tables),'docx_bullets':sum(p.style.name=='List Bullet' for p in doc.paragraphs),
    'english_runs_checked':count,'nonitalic_docx_runs':bad_docx,'nonitalic_pdf_spans':bad_pdf,
    'sentence_dashes_remaining':bool(re.search('[—–]',source)),
    'numbered_table_captions_remaining':bool(re.search(r'Табела \d+:',source)),
    'paragraphs_checked':len(prose),'missing_paragraphs':missing,
    'empty_pdf_pages':[i+1 for i,p in enumerate(pdf) if len(p.get_text().strip())<5]
}
(OUT/'provera_dokumenta.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
assert not bad_docx and not bad_pdf and not missing
assert len(doc.tables)==0 and report['docx_bullets']==sum(line.startswith('- ') for line in source.splitlines())
assert len(pdf)>=50
assert '{{SLIKA|' not in ''.join(page.get_text() for page in pdf)
assert not report['sentence_dashes_remaining'] and not report['numbered_table_captions_remaining'] and not report['empty_pdf_pages']
for idx,name in [(1,'controle_toc.png'),(10,'controle_bullets.png'),(13,'controle_page14.png')]:
    pdf[idx].get_pixmap().save(OUT/name)
