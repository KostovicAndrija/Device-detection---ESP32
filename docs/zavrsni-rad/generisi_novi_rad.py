from pathlib import Path
import re
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
SOURCE = OUT / 'Zavrsni_rad_ESP32.md'
ASSETS = OUT / 'assets_novi_rad'
ASSETS.mkdir(exist_ok=True)
DOCX = OUT / 'Zavrsni_rad_ESP32_prosireni.docx'

FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'
ITALIC = 'C:/Windows/Fonts/ariali.ttf'
BOLD_ITALIC = 'C:/Windows/Fonts/arialbi.ttf'

ENGLISH = re.compile(r'(?i)\b(?:ESP32|MQTT|RSSI|Wi-Fi|Bluetooth|Bluetooth Low Energy|BLE|Worker|API|REST|SignalR|Angular|PostgreSQL|Mosquitto|Docker|JSON|JWT|HTTP|GUID|QoS|AtLeastOnce|SQL|Entity Framework Core|BackgroundService|Dashboard|ingestion|pipeline|middleware|audit|retention|weighted centroid|weighted least squares|Kalman smoother|Problem Details|environment|frontend|backend|firmware|confidence|timestamp|eventId|deviceIdentifier|sensorId|sessionId|signalType|rssi|externalId|hash_id|first_seen|last_seen|type|student_ref|device_hash)\b')


def fnt(size, bold=False, italic=False):
    if bold and italic:
        path = BOLD_ITALIC
    elif bold:
        path = BOLD
    elif italic:
        path = ITALIC
    else:
        path = FONT
    return ImageFont.truetype(path, size)


def diagram(name, title, boxes, arrows):
    image = Image.new('RGB', (1700, 850), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((850, 38), title, font=fnt(34, True), fill='#1f3347', anchor='ma')
    for x, y, w, h, label in boxes:
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill='#eaf1f8', outline='#365b78', width=4)
        lines = label.split('\n')
        total = len(lines) * 36
        for i, line in enumerate(lines):
            draw.text((x + w / 2, y + h / 2 - total / 2 + i * 36), line, font=fnt(27, i == 0), fill='#1f3347', anchor='ma')
    for x1, y1, x2, y2, label in arrows:
        draw.line((x1, y1, x2, y2), fill='#47708e', width=5)
        angle = math.atan2(y2 - y1, x2 - x1)
        points = [(x2, y2), (x2 - 24 * math.cos(angle - .4), y2 - 24 * math.sin(angle - .4)), (x2 - 24 * math.cos(angle + .4), y2 - 24 * math.sin(angle + .4))]
        draw.polygon(points, fill='#47708e')
        if label:
            draw.text(((x1 + x2) / 2, (y1 + y2) / 2 - 24), label, font=fnt(22), fill='#47708e', anchor='ms')
    image.save(ASSETS / name, dpi=(240, 240))


def make_figures():
    diagram('arhitektura.png', 'Слојевита архитектура система', [
        (60, 170, 330, 130, 'ESP32\nсензори'), (500, 170, 330, 130, 'Mosquitto\nMQTT посредник'),
        (940, 170, 330, 130, '.NET Worker\nпријем порука'), (1380, 170, 250, 130, 'Application\nобрада'),
        (700, 500, 330, 130, 'PostgreSQL\nскладиште'), (1240, 500, 330, 130, 'ASP.NET Core\nAPI и SignalR'),
        (270, 500, 330, 130, 'Angular\nкориснички интерфејс')], [
        (390, 235, 500, 235, 'MQTT'), (830, 235, 940, 235, 'JSON'), (1270, 235, 1380, 235, 'услучивање'),
        (1100, 300, 865, 500, 'упис'), (1380, 300, 1400, 500, 'догађаји'), (1240, 565, 1030, 565, 'читање'), (700, 565, 600, 565, 'REST')])
    diagram('tok_obrade.png', 'Ток обраде једног RSSI опажања', [
        (70, 230, 280, 130, '1. Пријем\nMQTT поруке'), (400, 230, 280, 130, '2. Валидација\nи нормализација'),
        (730, 230, 280, 130, '3. Хеширање\nидентификатора'), (1060, 230, 280, 130, '4. Упис\nопажања'),
        (1390, 230, 240, 130, '5. Ризик\nи аларм'), (650, 550, 400, 130, '6. SignalR\nдогађај клијенту')], [
        (350, 295, 400, 295, ''), (680, 295, 730, 295, ''), (1010, 295, 1060, 295, ''), (1340, 295, 1390, 295, ''),
        (1510, 360, 1050, 550, 'резултат')])
    diagram('lokalizacija.png', 'Процена положаја у просторији', [
        (80, 180, 300, 130, 'Сензор S1\n(0, 0)'), (80, 500, 300, 130, 'Сензор S3\n(4, 6)'),
        (1320, 180, 300, 130, 'Сензор S2\n(8, 0)'), (600, 320, 500, 150, 'RSSI мерења\nи тежине'), (600, 570, 500, 150, 'WLS процена\nположаја')], [
        (380, 245, 600, 380, ''), (380, 565, 600, 640, ''), (1320, 245, 1100, 380, ''), (1100, 395, 1100, 640, '')])
    diagram('real_time.png', 'Путања података до корисничког интерфејса', [
        (70, 250, 300, 130, 'Сензорско\nочитавање'), (470, 250, 300, 130, 'Ingestion\nсервис'),
        (870, 250, 300, 130, 'База и\nпословна логика'), (1270, 250, 300, 130, 'Angular\nприказ'),
        (870, 560, 300, 130, 'SignalR\nдогађај')], [
        (370, 315, 470, 315, 'HTTP/MQTT'), (770, 315, 870, 315, 'упис'), (1170, 315, 1270, 315, 'REST'), (1020, 380, 1020, 560, 'push')])


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    add_runs(p, text, size=9.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_runs(paragraph, text, size=12, bold=False):
    position = 0
    for match in ENGLISH.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(size)
            run.bold = bold
        run = paragraph.add_run(match.group(0))
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold


def add_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(8)
    add_runs(p, text)
    return p


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, value in enumerate(rows[0]):
        set_cell_text(table.rows[0].cells[i], value, True)
        set_cell_shading(table.rows[0].cells[i], 'DCE6F1')
    header = table.rows[0]._tr.get_or_add_trPr()
    header.append(OxmlElement('w:tblHeader'))
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
        trpr = table.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, end])


def add_toc(doc, headings):
    doc.add_heading('Садржај', level=1)
    for level, title in headings:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm((level - 1) * 0.6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        add_runs(p, title, size=11)


def make_doc():
    make_figures()
    text = SOURCE.read_text(encoding='utf-8').splitlines()
    headings = []
    for line in text:
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            if level <= 2 and not line[level:].strip().startswith('9.'):
                headings.append((level, line[level:].strip()))

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    add_page_number(section)

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in [('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(15)
        style.paragraph_format.space_after = Pt(9)
        style.paragraph_format.keep_with_next = True
        if name == 'Heading 1':
            style.paragraph_format.page_break_before = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(25)
    p.add_run('ИНСТИТУТ ЗА МАТЕМАТИКУ И ИНФОРМАТИКУ\n').bold = True
    p.add_run('ПРИРОДНО МАТЕМАТИЧКИ ФАКУЛТЕТ\n').bold = True
    p.add_run('УНИВЕРЗИТЕТ У КРАГУЈЕВЦУ').bold = True
    p.paragraph_format.line_spacing = 1.35
    for r in p.runs:
        r.font.name = 'Times New Roman'; r.font.size = Pt(12)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('ЗАВРШНИ РАД')
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(16)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.2
    add_runs(p, 'СИСТЕМ ЗА ДЕТЕКЦИЈУ БЕЖИЧНИХ\nУРЕЂАЈА ПРИМЕНОМ ESP32 СЕНЗОРА', size=18, bold=True)
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.right_indent = Cm(1)
    p.paragraph_format.line_spacing = 1.4
    add_runs(p, 'Ментор: др Ана Капларевић Малишић\nСтудент: Андрија Костовић, 51/2021', size=12, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph('Крагујевац, 2026.')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.name = 'Times New Roman'; p.runs[0].font.size = Pt(12)

    doc.add_page_break()
    for line in text:
        if line.startswith('# 1.'):
            break
    add_toc(doc, headings)
    doc.add_page_break()

    i = 0
    skip_section = False
    skip_code = False
    figure_map = {
        'Слојевита архитектура': 'arhitektura.png',
        'Ток обраде': 'tok_obrade.png',
        'Процена положаја': 'lokalizacija.png',
        'Путања података': 'real_time.png'}
    while i < len(text):
        line = text[i].strip()
        if not line:
            i += 1
            continue
        if re.search(r'(?i)симулатор|simulator', line):
            i += 1
            continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            title = line[level:].strip()
            if level == 1 and title.startswith('9. Тестирање'):
                skip_section = True
                i += 1
                continue
            if level == 1:
                skip_section = False
            if skip_section:
                i += 1
                continue
            if title in ('Сажетак', 'Abstract'):
                doc.add_heading(title, level=1)
            else:
                doc.add_heading(title, level=level)
            i += 1
            continue
        if skip_section:
            i += 1
            continue
        if line.startswith('```'):
            skip_code = not skip_code
            i += 1
            continue
        if skip_code:
            i += 1
            continue
        if line.startswith('![(') or line.startswith('!['):
            match = re.match(r'!\[[^\]]*\]\(([^)]+)\)', line)
            if match:
                image_path = OUT / match.group(1)
                if not image_path.exists():
                    image_path = OUT / 'assets' / Path(match.group(1)).name
                if image_path.exists():
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.add_run().add_picture(str(image_path), width=Cm(15.5))
                    caption = re.match(r'!\[([^\]]*)\]', line).group(1)
                    doc.add_paragraph(caption, style='Caption')
                i += 1
                continue
        if line.startswith('{{SLIKA|'):
            parts = line.strip('{}').split('|')
            if len(parts) >= 3:
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(paragraph, f'[Место за слику: {parts[2]}]', size=10)
                paragraph.paragraph_format.space_before = Pt(5)
                paragraph.paragraph_format.space_after = Pt(5)
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(text) and text[i].strip().startswith('|'):
                parts = [v.strip() for v in text[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+', p) for p in parts):
                    rows.append(parts)
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith('$$'):
            formula = line.strip('$')
            i += 1
            while i < len(text) and not text[i].strip().startswith('$$'):
                formula += ' ' + text[i].strip(); i += 1
            i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(9)
            r = p.add_run(formula)
            r.font.name = 'Cambria Math'; r.font.size = Pt(12); r.italic = True
            continue
        if line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. ') or line.startswith('5. ') or line.startswith('6. ') or line.startswith('7. ') or line.startswith('8. ') or line.startswith('9. ') or line.startswith('10. ') or line.startswith('11. '):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.line_spacing = 1.2; p.paragraph_format.space_after = Pt(4)
            add_runs(p, re.sub(r'^\d+\. ', '', line))
            i += 1
            continue
        if line.startswith('**') and line.endswith('**'):
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
            add_runs(p, line.strip('*'), bold=True)
            i += 1
            continue
        p = add_paragraph(doc, line)
        for key, filename in figure_map.items():
            if key in line and i + 1 < len(text):
                q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                q.paragraph_format.space_before = Pt(5); q.paragraph_format.space_after = Pt(3)
                q.add_run().add_picture(str(ASSETS / filename), width=Cm(15.5))
                caption = doc.add_paragraph('Слика. ' + key + '.', style='Caption')
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                break
        i += 1

    doc.save(DOCX)
    print(DOCX)


if __name__ == '__main__':
    make_doc()
